"""Read a teacher roster from an Excel (.xlsx) or CSV file.

Pure parsing and shape validation -- nothing here touches MongoDB. Whether an
address already has an account is decided by the router, so these functions
stay testable without a database (see the api skill on validators).

The file is whatever a department office has to hand, so the reader is
forgiving about layout: the header row may use any of the common names for
each column, columns may come in any order, extra columns are ignored, and a
file with no header at all is read by recognising which column holds email
addresses and which holds phone numbers.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from typing import Any

from pydantic import EmailStr, TypeAdapter, ValidationError

from app.schemas.common import normalize_phone


# A roster of a few hundred teachers is a few tens of kilobytes even as .xlsx.
MAX_FILE_BYTES = 2 * 1024 * 1024
# Rows accepted from one file. Matches the superadmin's bulk onboarding cap.
MAX_ROWS = 500
# Rows read before giving up on a sheet, blank ones included, so a sheet with
# a million formatted-but-empty rows cannot keep the request busy.
MAX_SCAN_ROWS = 5000

_XLSX_MAGIC = b"PK\x03\x04"
_XLS_MAGIC = b"\xd0\xcf\x11\xe0"  # the pre-2007 binary .xls (OLE) format

COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "name": (
        "name", "full name", "teacher name", "teacher", "faculty name", "faculty",
        "employee name", "staff name",
    ),
    "email": (
        "email", "e mail", "email id", "e mail id", "email address", "mail",
        "mail id", "official email", "official email id", "username",
    ),
    "phone": (
        "phone", "mobile", "mobile number", "mobile no", "phone number", "phone no",
        "contact", "contact number", "contact no", "cell", "whatsapp", "whatsapp number",
    ),
    "department": ("department", "dept", "department name", "dept name", "school"),
    "designation": ("designation", "title", "rank", "position", "post"),
}

_email_adapter = TypeAdapter(EmailStr)


class ImportFileError(ValueError):
    """The file cannot be read as a roster. The message is safe to show."""


# ============================================================
# FILE -> CELLS
# ============================================================

def _cell_text(value: Any) -> str:
    """One spreadsheet cell as text, the way the person typed it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    # Excel stores a typed mobile number as a float: 9876543210.0.
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return " ".join(str(value).split())


def _read_xlsx(content: bytes) -> list[list[str]]:
    # Imported lazily: only this path needs it.
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise ImportFileError(
            "This Excel file could not be read. Open it in Excel and save it again as .xlsx."
        ) from error

    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ImportFileError("This Excel file has no worksheets.")
        rows: list[list[str]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= MAX_SCAN_ROWS:
                break
            rows.append([_cell_text(value) for value in row])
        return rows
    finally:
        workbook.close()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[str]] = []
    for index, row in enumerate(reader):
        if index >= MAX_SCAN_ROWS:
            break
        rows.append([_cell_text(cell) for cell in row])
    return rows


def read_rows(filename: str | None, content: bytes) -> list[list[str]]:
    """The file's cells as text rows. Raises ``ImportFileError``."""
    if not content:
        raise ImportFileError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(
            f"The file is too large. The limit is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )

    # Decided by the bytes, not the name: a renamed file is still read right.
    if content.startswith(_XLS_MAGIC):
        raise ImportFileError(
            "Old .xls files are not supported. Open the file in Excel and save it as .xlsx."
        )
    if content.startswith(_XLSX_MAGIC):
        return _read_xlsx(content)

    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        raise ImportFileError("This does not look like a valid Excel file.")
    return _read_csv(content)


# ============================================================
# CELLS -> ROSTER ROWS
# ============================================================

def _header_key(text: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text.lower()).split())


_ALIAS_LOOKUP = {
    alias: field for field, aliases in COLUMN_ALIASES.items() for alias in aliases
}


def _looks_like_email(text: str) -> bool:
    return "@" in text and "." in text.rsplit("@", 1)[-1]


def _looks_like_phone(text: str) -> bool:
    return len(re.sub(r"\D", "", text)) >= 10 and not _looks_like_email(text)


# Header words that mark a column as a counter or an identifier, never a name:
# "S.No", "Sr. No.", "Employee ID", "Emp Code".
_NOT_A_NAME = {"s", "sr", "sl", "sno", "srno", "slno", "serial", "no", "number", "num",
               "id", "code", "roll", "index"}
# Somebody else's name.
_OTHER_PERSON = {"father", "mother", "husband", "spouse", "guardian", "hod", "reporting"}


def _field_for_header(cell: str) -> str | None:
    """Which roster field a header cell names, or None.

    Exact aliases first; then by keyword, because real sheets say "Official
    E-mail", "Name of the Faculty" or "Contact No. (Mobile)" as often as the
    plain words. A cell holding an email address is data, never a header.
    """
    if not cell or _looks_like_email(cell):
        return None
    key = _header_key(cell)
    if not key:
        return None
    if key in _ALIAS_LOOKUP:
        return _ALIAS_LOOKUP[key]

    words = set(key.split())
    if "email" in words or "mail" in words or key.startswith("e mail"):
        return "email"
    if words & {"mobile", "phone", "contact", "whatsapp", "cell", "telephone", "tel"}:
        return "phone"
    if words & {"designation", "rank", "position", "post"}:
        return "designation"
    if words & {"department", "dept", "school", "discipline"}:
        return "department"
    if "name" in words and not words & _OTHER_PERSON:
        return "name"
    if words & {"teacher", "faculty", "employee", "staff"} and not words & (_NOT_A_NAME | _OTHER_PERSON):
        return "name"
    return None


def _has_letters(text: str) -> bool:
    return bool(re.search(r"[^\W\d_]", text))


def _locate_columns(rows: list[list[str]]) -> tuple[dict[str, int], int]:
    """Map each known field to a column index; return it with the first data row.

    The header is the first of the opening rows that names an email column, so
    a title line ("SRHU Faculty List 2026") or a blank row above it is fine.
    Columns it does not recognise -- a serial number, an employee ID -- are
    simply not used. Only when no row names an email column is the layout
    worked out from the cell contents instead.
    """
    for header_index, row in enumerate(rows[:15]):
        columns: dict[str, int] = {}
        for index, cell in enumerate(row):
            field = _field_for_header(cell)
            if field and field not in columns:
                columns[field] = index
        if "email" in columns:
            return columns, header_index + 1

    # No header: recognise the columns by what they hold.
    sample = [row for row in rows[:50] if any(row)]
    width = max((len(row) for row in sample), default=0)

    def share(index: int, test) -> float:
        cells = [row[index] for row in sample if index < len(row) and row[index]]
        return sum(1 for cell in cells if test(cell)) / len(cells) if cells else 0.0

    columns = {}
    email_scores = [(share(i, _looks_like_email), i) for i in range(width)]
    best_email = max(email_scores, default=(0.0, -1))
    if best_email[0] < 0.5:
        raise ImportFileError(
            "No email column was found. Add a header row with an \"Email\" column."
        )
    columns["email"] = best_email[1]

    phone_scores = [(share(i, _looks_like_phone), i) for i in range(width) if i != columns["email"]]
    best_phone = max(phone_scores, default=(0.0, -1))
    if best_phone[0] >= 0.5:
        columns["phone"] = best_phone[1]

    # A name is words. A column of serial numbers or IDs is never taken for
    # one; with no column of words, names are made from the email addresses.
    def is_name(cell: str) -> bool:
        return _has_letters(cell) and not _looks_like_email(cell) and not _looks_like_phone(cell)

    for i in range(width):
        if i not in columns.values() and share(i, is_name) >= 0.8:
            columns["name"] = i
            break
    return columns, 0


def name_from_email(email: str) -> str:
    """A readable placeholder name when the file gives none: rajesh.kumar -> Rajesh Kumar."""
    local = email.split("@", 1)[0]
    words = [w for w in re.split(r"[._\-+]+|\d+", local) if w]
    return " ".join(w.capitalize() for w in words) or local


def _valid_email(value: str) -> bool:
    try:
        _email_adapter.validate_python(value)
    except ValidationError:
        return False
    return True


def parse_roster(rows: list[list[str]]) -> tuple[list[dict], bool]:
    """Turn cell rows into roster entries; return them and whether rows were cut off.

    Each entry: ``row`` (the spreadsheet row number), ``name``, ``email``,
    ``phone``, ``department``, ``designation``, ``name_derived``, ``status``
    (``"ok"``, ``"duplicate"`` or ``"invalid"``), ``reason`` and ``warnings``.
    """
    columns, start = _locate_columns(rows)

    def cell(row: list[str], field: str) -> str:
        index = columns.get(field)
        return row[index].strip() if index is not None and index < len(row) else ""

    entries: list[dict] = []
    seen: dict[str, int] = {}
    truncated = False

    for offset, row in enumerate(rows[start:]):
        if not any(cell_value.strip() for cell_value in row):
            continue
        if len(entries) >= MAX_ROWS:
            truncated = True
            break

        email = cell(row, "email").casefold()
        entry: dict[str, Any] = {
            "row": start + offset + 1,
            "name": " ".join(cell(row, "name").split())[:120] or None,
            "email": email,
            "phone": None,
            "department": cell(row, "department")[:120] or None,
            "designation": cell(row, "designation")[:120] or None,
            "name_derived": False,
            "status": "ok",
            "reason": None,
            "warnings": [],
        }

        raw_phone = cell(row, "phone")
        if raw_phone:
            try:
                entry["phone"] = normalize_phone(raw_phone)
            except ValueError:
                entry["warnings"].append(
                    f"Mobile number \"{raw_phone}\" is not 10 digits and will be left blank."
                )

        if not email:
            entry.update(status="invalid", reason="Email address is missing.")
        elif not _valid_email(email):
            entry.update(status="invalid", reason="Email address is not valid.")
        elif email in seen:
            entry.update(status="duplicate", reason=f"Same email as row {seen[email]}.")
        else:
            seen[email] = entry["row"]

        if entry["status"] == "ok" and not entry["name"]:
            entry["name"] = name_from_email(email)
            entry["name_derived"] = True

        entries.append(entry)

    return entries, truncated
